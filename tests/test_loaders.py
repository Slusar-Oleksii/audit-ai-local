from pathlib import Path

from docx import Document

from audit_ai.loaders.csv_loader import load_csv
from audit_ai.loaders.docx import load_docx
from audit_ai.loaders.pdf import _needs_ocr
from audit_ai.loaders.text_code import load_text_or_code
from audit_ai.schemas import TextUnit


def test_code_loader_preserves_line_numbers(tmp_path: Path):
    path = tmp_path / "sample.py"
    path.write_text("x = 1\n\nprint(x)\n", encoding="utf-8")
    units = load_text_or_code(path, is_code=True)
    assert [unit.line_start for unit in units] == [1, 2, 3]
    assert units[1].text == ""
    assert all(unit.kind == "code" for unit in units)


def test_csv_loader_detects_semicolon(tmp_path: Path):
    path = tmp_path / "report.csv"
    path.write_text("name;amount\nAlpha;10\nBeta;20\n", encoding="utf-8")
    units = load_csv(path)
    assert units[0].row_start == 2
    assert "name: Alpha" in units[0].text
    assert "amount: 20" in units[1].text


def test_csv_loader_preserves_columns_missing_from_header(tmp_path: Path):
    path = tmp_path / "broken.csv"
    path.write_text("name,amount\nAlpha,10,unexpected\n", encoding="utf-8")
    units = load_csv(path)
    assert "extra_column_1: unexpected" in units[0].text


def test_windows_1251_text_and_csv_are_supported(tmp_path: Path):
    text_path = tmp_path / "notes.txt"
    text_path.write_bytes("Фінансовий звіт".encode("cp1251"))
    assert "Фінансовий" in load_text_or_code(text_path, is_code=False)[0].text

    csv_path = tmp_path / "report.csv"
    csv_path.write_bytes("назва;сума\nПослуги;100\n".encode("cp1251"))
    units = load_csv(csv_path)
    assert "назва: Послуги" in units[0].text


def test_docx_loader_reads_paragraphs_and_tables(tmp_path: Path, settings):
    path = tmp_path / "contract.docx"
    document = Document()
    document.add_paragraph("Умови оплати — 10 днів.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Стаття"
    table.cell(0, 1).text = "Сума"
    table.cell(1, 0).text = "Послуги"
    table.cell(1, 1).text = "1000"
    document.add_paragraph("Після таблиці")
    document.save(path)

    units = load_docx(path, settings)
    assert any("Умови оплати" in unit.text for unit in units)
    assert any("Послуги | 1000" in unit.text for unit in units)
    assert [unit.kind for unit in units] == ["text", "table", "text"]
    assert units[-1].text == "Після таблиці"


def test_pdf_ocr_detection():
    assert _needs_ocr([TextUnit(text="", page=1), TextUnit(text="", page=2)])
    assert not _needs_ocr([TextUnit(text="Достатньо тексту " * 20, page=1)])
